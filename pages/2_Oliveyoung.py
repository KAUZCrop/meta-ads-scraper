# ============================================================
# 올리브영 리뷰 수집기 — 별도 페이지
# ============================================================
import asyncio
import concurrent.futures
import io
import json
import os
import queue
import random
import re
import sys
import time
from typing import Callable

import pandas as pd
import streamlit as st

# ── CSS 공유 ─────────────────────────────────────────────────
_css_path = os.path.join(os.path.dirname(__file__), "..", "static", "style.css")
try:
    with open(_css_path) as _f:
        st.markdown(f"<style>{_f.read()}</style>", unsafe_allow_html=True)
except FileNotFoundError:
    pass

st.set_page_config(
    page_title="올리브영 리뷰 — ADINTEL",
    page_icon="🌿",
    layout="wide",
)

# ============================================================
# 올리브영 리뷰 스크래퍼 (self-contained)
# ============================================================
_DETAIL_URL    = "https://www.oliveyoung.co.kr/store/goods/getGoodsDetail.do"
_HOME_URL      = "https://www.oliveyoung.co.kr"
_CHROMIUM_PATH = "/usr/bin/chromium"

_USER_AGENTS = [
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/136.0.0.0 Safari/537.36",
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/136.0.0.0 Safari/537.36",
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/135.0.0.0 Safari/537.36",
    "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/135.0.0.0 Safari/537.36",
]

_ANTI_DETECT_JS = """
    Object.defineProperty(navigator, 'webdriver', { get: () => undefined });
    Object.defineProperty(navigator, 'plugins',   { get: () => [1, 2, 3, 4, 5] });
    window.chrome = { runtime: {} };
"""

_EXTRACT_JS = """
() => {
    function findInShadow(root, tagName) {
        const found = Array.from(root.querySelectorAll(tagName));
        for (const el of root.querySelectorAll('*')) {
            if (el.shadowRoot) {
                const nested = findInShadow(el.shadowRoot, tagName);
                for (const n of nested) found.push(n);
            }
        }
        return found;
    }
    const container = document.querySelector('oy-review-review-in-product');
    if (!container?.shadowRoot) return { error: 'no_container', reviews: [] };
    const items = findInShadow(container.shadowRoot, 'oy-review-review-item');
    if (!items.length) return { error: 'no_items', reviews: [] };
    const reviews = [];
    for (const item of items) {
        if (!item.shadowRoot) continue;
        const sr = item.shadowRoot;
        const userComp = sr.querySelector('oy-review-review-user');
        const name     = userComp?.shadowRoot?.querySelector('.name')?.textContent?.trim() || '';
        const skinType = userComp?.shadowRoot?.querySelector('.skin-type')?.textContent?.trim() || '';
        const date    = sr.querySelector('.date')?.textContent?.trim() || '';
        const option  = sr.querySelector('.goods-option')?.textContent?.trim() || '';
        const contentComp = sr.querySelector('oy-review-review-content');
        const content = contentComp?.shadowRoot?.querySelector('p')?.textContent?.trim() || '';
        const helpful = sr.querySelector('.helpful-count')?.textContent?.trim() || '0';
        if (name || content) {
            reviews.push({ name, date, option, content, skinType, helpful });
        }
    }
    return { reviews, count: reviews.length };
}
"""

_WAIT_READY_JS = "() => { const el = document.querySelector('oy-review-review-in-product'); return !!(el && el.shadowRoot); }"

_WAIT_ITEMS_JS = """
() => {
    const el = document.querySelector('oy-review-review-in-product');
    if (!el?.shadowRoot) return false;
    for (const node of el.shadowRoot.querySelectorAll('*')) {
        if (node.tagName?.toLowerCase() === 'oy-review-review-item') return true;
        if (node.shadowRoot) {
            if (node.shadowRoot.querySelector('oy-review-review-item')) return true;
        }
    }
    return false;
}
"""


def _to_review(r: dict) -> dict:
    return {
        "reviewer":      r.get("name", "익명"),
        "date":          r.get("date", ""),
        "title":         "",
        "content":       r.get("content", ""),
        "skin_type":     r.get("skinType", ""),
        "helpful":       r.get("helpful", "0"),
        "purchase_type": r.get("option", ""),
    }


async def _run_scrape(goods_no: str, max_pages: int, on_progress: Callable, log: list) -> list:
    from playwright.async_api import async_playwright, TimeoutError as PWTimeout

    executable = _CHROMIUM_PATH if os.path.exists(_CHROMIUM_PATH) else None
    log.append(f"[시작] goodsNo={goods_no}, 최대 {max_pages}페이지")

    all_reviews: list = []
    seen: set = set()

    async with async_playwright() as pw:
        launch_kwargs = dict(
            headless=True,
            args=[
                "--disable-blink-features=AutomationControlled",
                "--no-sandbox", "--disable-dev-shm-usage", "--disable-gpu",
            ],
        )
        if executable:
            launch_kwargs["executable_path"] = executable

        browser = await pw.chromium.launch(**launch_kwargs)
        context = await browser.new_context(
            viewport={"width": 1920, "height": 1080},
            user_agent=random.choice(_USER_AGENTS),
            locale="ko-KR",
        )
        await context.add_init_script(_ANTI_DETECT_JS)
        page = await context.new_page()

        try:
            # ── 1단계: 홈 방문 → 쿠키 확보 ──────────────────────
            log.append("[1단계] 올리브영 홈 방문")
            await page.goto(_HOME_URL, wait_until="domcontentloaded", timeout=30000)
            await asyncio.sleep(random.uniform(1.5, 2.5))

            # ── 2단계: 상품 상세 페이지 이동 ──────────────────────
            detail_url = f"{_DETAIL_URL}?goodsNo={goods_no}"
            log.append(f"[2단계] 상품 페이지 이동: {detail_url}")
            await page.goto(detail_url, wait_until="domcontentloaded", timeout=30000)
            await asyncio.sleep(random.uniform(2.0, 3.0))

            # 팝업 닫기
            for selector in ["button.btn_close", ".layer_close", ".btnClose", "button[aria-label='닫기']"]:
                try:
                    btn = page.locator(selector).first
                    if await btn.is_visible(timeout=500):
                        await btn.click()
                        await asyncio.sleep(0.3)
                except Exception:
                    pass

            # ── 3단계: 리뷰 탭 클릭 ─────────────────────────────
            log.append("[3단계] 리뷰 탭 클릭 시도")
            for tab_sel in [
                'a[href="#reviewInfo"]', 'button:has-text("리뷰")',
                'a:has-text("리뷰")', '#reviewInfo',
            ]:
                try:
                    el = page.locator(tab_sel).first
                    if await el.is_visible(timeout=1500):
                        await el.click()
                        await asyncio.sleep(1.5)
                        log.append(f"[3단계] 탭 클릭 성공: {tab_sel}")
                        break
                except Exception:
                    pass

            # ── 4단계: 리뷰 섹션으로 스크롤 ──────────────────────
            log.append("[4단계] 리뷰 섹션 스크롤")
            for anchor in ["#reviewInfo", "oy-review-review-in-product"]:
                try:
                    await page.evaluate(f"""
                        () => {{
                            const el = document.querySelector('{anchor}');
                            if (el) el.scrollIntoView({{behavior:'smooth', block:'center'}});
                        }}
                    """)
                    await asyncio.sleep(1.0)
                except Exception:
                    pass

            # ── 5단계: Shadow DOM 컨테이너 대기 ───────────────────
            log.append("[5단계] Shadow DOM 컨테이너 대기")
            try:
                await page.wait_for_function(_WAIT_READY_JS, timeout=15000)
                log.append("[5단계] Shadow DOM 컨테이너 감지됨")
                await page.wait_for_function(_WAIT_ITEMS_JS, timeout=10000)
                log.append("[5단계] 리뷰 아이템 감지됨")
            except PWTimeout:
                log.append("[5단계] 타임아웃 — 리뷰 로드 실패")

            await asyncio.sleep(1.0)

            # ── 6단계: 추출 → 스크롤 반복 ────────────────────────
            target = max_pages * 10
            no_new_streak = 0
            scroll_num = 0
            log.append(f"[6단계] 추출→스크롤 반복 (목표 {target}개)")

            while len(all_reviews) < target:
                result = await page.evaluate(_EXTRACT_JS)
                raw = result.get("reviews", [])
                new_batch = [_to_review(r) for r in raw
                             if r.get("content") not in seen]

                if new_batch:
                    all_reviews.extend(new_batch)
                    seen.update(r["content"] for r in new_batch if r.get("content"))
                    no_new_streak = 0
                    log.append(f"[스크롤{scroll_num}] +{len(new_batch)}개 (누적 {len(all_reviews)}개)")
                else:
                    no_new_streak += 1
                    if no_new_streak >= 4:
                        log.append(f"[스크롤{scroll_num}] 4회 연속 신규 없음 → 종료")
                        break

                if on_progress:
                    on_progress(
                        min(len(all_reviews) // 10 + 1, max_pages),
                        max_pages,
                        len(all_reviews),
                    )

                scroll_num += 1
                await page.evaluate("window.scrollBy(0, 300)")
                await asyncio.sleep(0.6)

        except Exception as e:
            log.append(f"[오류] {type(e).__name__}: {e}")
        finally:
            await browser.close()

    log.append(f"[완료] 최종 수집: {len(all_reviews)}개")
    return all_reviews


def _run_in_thread(goods_no: str, max_pages: int, on_progress: Callable, log: list) -> list:
    return asyncio.run(_run_scrape(goods_no, max_pages, on_progress, log))


def scrape_reviews(
    goods_no: str,
    max_pages: int = 10,
    on_progress: Callable = None,
) -> tuple[list, list]:
    """(리뷰 리스트, 진단 로그) 반환"""
    log: list = []
    progress_q: queue.Queue = queue.Queue()

    def _thread_progress(cur, total, collected):
        progress_q.put((cur, total, collected))

    with concurrent.futures.ThreadPoolExecutor(max_workers=1) as executor:
        future = executor.submit(
            _run_in_thread, goods_no, max_pages,
            _thread_progress if on_progress else None,
            log,
        )
        if on_progress:
            while not future.done():
                try:
                    on_progress(*progress_q.get(timeout=0.1))
                except queue.Empty:
                    pass
            while True:
                try:
                    on_progress(*progress_q.get_nowait())
                except queue.Empty:
                    break
        return future.result(), log


# ============================================================
# 유틸 함수
# ============================================================
def _extract_goods_no(text: str) -> str:
    text = text.strip()
    if re.match(r'^A\d{9,}$', text):
        return text
    match = re.search(r'goodsNo=([A-Za-z0-9]+)', text)
    return match.group(1) if match else ""


def _stars(rating) -> str:
    try:
        r = max(0, min(5, int(float(rating))))
    except Exception:
        return "☆☆☆☆☆"
    return "★" * r + "☆" * (5 - r)


def _to_csv(df: pd.DataFrame) -> bytes:
    buf = io.BytesIO()
    df.to_csv(buf, index=False, encoding="utf-8-sig")
    return buf.getvalue()


def _to_excel(df: pd.DataFrame, pname: str = "", reviews: list = None) -> bytes:  # noqa: ARG001
    try:
        import openpyxl
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        return _to_csv(df)

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        df.to_excel(w, index=False, sheet_name="리뷰")
        ws = w.sheets["리뷰"]
        for col, width in zip("ABCDEFGH", [14, 14, 24, 60, 12, 10, 14, 14]):
            ws.column_dimensions[col].width = width
        fill = PatternFill(start_color="1A5C38", end_color="1A5C38", fill_type="solid")
        for cell in ws[1]:
            cell.fill = fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    return buf.getvalue()


def _ko_df(df: pd.DataFrame) -> pd.DataFrame:
    return df.rename(columns={
        "reviewer": "작성자", "date": "작성일",
        "title": "제목", "content": "내용", "skin_type": "피부타입",
        "helpful": "도움돼요", "purchase_type": "구매유형",
    })


# ============================================================
# 페이지 UI
# ============================================================
# 헤더
st.markdown(
    '<div class="hdr">'
    '<div class="hdr-row">'
    '<div class="logo">AD<b>INTEL</b></div>'
    '<span class="ver">올리브영 리뷰</span>'
    '</div>'
    '<div class="sub">OLIVEYOUNG REVIEW SCRAPER · 리뷰 수집 도구</div>'
    '</div>',
    unsafe_allow_html=True,
)

# ── 세션 초기화 ──────────────────────────────────────────────
for _k, _v in {"oy_product": None, "oy_reviews": []}.items():
    if _k not in st.session_state:
        st.session_state[_k] = _v

# ── 사이드바 ────────────────────────────────────────────────
with st.sidebar:
    st.markdown('<div class="slbl">상품 입력</div>', unsafe_allow_html=True)

    with st.form(key="oy_form"):
        url_input  = st.text_area(
            "상품 URL 또는 goodsNo",
            placeholder=(
                "https://www.oliveyoung.co.kr/store/goods/getGoodsDetail.do?goodsNo=A000000200106\n"
                "또는 goodsNo만: A000000200106"
            ),
            height=120,
            label_visibility="collapsed",
        )
        pname_input = st.text_input("상품명 (파일명용)", placeholder="예: 투에이엔 선크림")
        max_pages   = st.slider("수집 페이지 수", 1, 50, 10, help="1페이지 ≈ 리뷰 10개")
        submit_btn  = st.form_submit_button("🌿 리뷰 수집 시작", use_container_width=True)

    st.markdown("---")
    if st.button("🔄 초기화", use_container_width=True):
        for _k in ["oy_product", "oy_reviews"]:
            st.session_state[_k] = None if _k == "oy_product" else []
        st.rerun()

    st.markdown("---")
    st.markdown(
        '<div style="font-size:10px;color:var(--mu);">⚠️ 개인·연구 목적으로만 사용하세요.</div>',
        unsafe_allow_html=True,
    )

# ── 수집 실행 ───────────────────────────────────────────────
if submit_btn:
    goods_no = _extract_goods_no(url_input or "")
    if not goods_no:
        st.error("URL 또는 goodsNo를 올바르게 입력해주세요.")
    else:
        pname = (pname_input or "").strip() or goods_no
        st.session_state.oy_product = {"name": pname, "goods_no": goods_no}
        st.session_state.oy_reviews = []

        st.markdown(
            f'<div class="banner">'
            f'<div class="dot dot-on"></div>'
            f'<span class="banner-txt">인식된 goodsNo: <b>{goods_no}</b> · 상품명: {pname}</span>'
            f'</div>',
            unsafe_allow_html=True,
        )

        pbar   = st.progress(0, text="수집 준비 중...")
        status = st.empty()

        def _on_progress(cur, total, collected):
            pbar.progress(int(cur / total * 100), text=f"페이지 {cur} / {total} 수집 중...")
            status.caption(f"📦 현재까지 {collected}개 리뷰 수집됨")

        with st.spinner("Playwright 브라우저로 리뷰 수집 중..."):
            reviews, diag_log = scrape_reviews(goods_no, max_pages=max_pages, on_progress=_on_progress)

        pbar.progress(100, text="✅ 완료!")
        status.empty()

        if not reviews:
            st.error("리뷰를 가져오지 못했습니다.")
            with st.expander("🔍 진단 로그", expanded=True):
                for line in diag_log:
                    st.text(line)
            st.session_state.oy_product = None
        else:
            st.session_state.oy_reviews = reviews
            st.rerun()

# ── 결과 표시 ───────────────────────────────────────────────
if st.session_state.oy_reviews:
    reviews = st.session_state.oy_reviews
    pname   = (st.session_state.oy_product or {}).get("name", "상품")[:30]
    df      = _ko_df(pd.DataFrame(reviews))

    # KPI
    k1, k2 = st.columns(2)
    k1.metric("총 리뷰 수",  f"{len(reviews):,}개")
    k2.metric("수집 페이지", f"{(len(reviews) - 1) // 10 + 1}페이지")

    st.markdown("---")

    # 다운로드
    st.markdown(
        '<div class="sec"><div class="sec-t">다운로드</div></div>',
        unsafe_allow_html=True,
    )
    d1, d2, d3, _ = st.columns([1, 1, 1, 3])
    with d1:
        st.download_button(
            "📄 CSV", data=_to_csv(df),
            file_name=f"{pname}_리뷰_{time.strftime('%Y%m%d')}.csv",
            mime="text/csv", use_container_width=True,
        )
    with d2:
        st.download_button(
            "📊 Excel", data=_to_excel(df, pname=pname, reviews=reviews),
            file_name=f"{pname}_리뷰_{time.strftime('%Y%m%d')}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )
    with d3:
        st.download_button(
            "🗂 JSON",
            data=json.dumps(reviews, ensure_ascii=False, indent=2).encode("utf-8"),
            file_name=f"{pname}_리뷰_{time.strftime('%Y%m%d')}.json",
            mime="application/json", use_container_width=True,
        )

    st.markdown("---")

    # 리뷰 / 테이블 탭
    tab_list, tab_table = st.tabs(["💬 리뷰 목록", "📋 데이터 테이블"])

    with tab_list:
        st.caption(f"{len(reviews)}개")
        for r in reviews[:100]:
            skin    = f" · {r['skin_type']}" if r.get("skin_type") else ""
            helpful = f" · 도움돼요 {r['helpful']}" if r.get("helpful") else ""
            content = str(r.get("content", "")).replace("\n", "<br>")
            st.markdown(
                f'<div class="card">'
                f'<div class="card-body">'
                f'<div style="display:flex;justify-content:space-between;align-items:center;">'
                f'<span style="font-weight:700;font-size:13px;">{r.get("reviewer","익명")}</span>'
                f'<span style="font-size:11px;color:var(--mu);">{r.get("date","")}</span>'
                f'</div>'
                f'<div class="card-meta">{skin}{helpful}</div>'
                f'<div class="card-cap">{content}</div>'
                f'</div></div>',
                unsafe_allow_html=True,
            )
        if len(reviews) > 100:
            st.info(f"100개까지 표시됩니다. 전체 {len(reviews)}개는 다운로드로 확인하세요.")

    with tab_table:
        st.dataframe(df, use_container_width=True, height=520)

else:
    # 빈 상태 안내
    st.markdown(
        '<div class="empty">'
        '<div class="empty-t">🌿 올리브영 리뷰 수집기</div>'
        '<div class="empty-d">'
        '왼쪽 사이드바에 상품 URL 또는 goodsNo를 입력하고<br>'
        '수집 버튼을 누르면 리뷰를 자동으로 가져옵니다.'
        '</div></div>',
        unsafe_allow_html=True,
    )
